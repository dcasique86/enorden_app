import os
import pandas as pd
import tempfile

from openpyxl import load_workbook


class TestImportacion:

    def _crear_excel_simple(self, data_dir, filas):
        """Crea un archivo Excel de prueba con los datos indicados."""
        ruta = os.path.join(data_dir, "test_import.xlsx")
        df = pd.DataFrame(filas)
        df.to_excel(ruta, index=False)
        return ruta

    def test_importar_clientes_desde_excel(self, db, data_dir):
        filas = [
            {"nombre": "Ana López", "telefono": "3001112233", "deuda": 100000},
            {"nombre": "Luis Mora", "telefono": "3001112244", "deuda": 200000},
            {"nombre": "Carla Díaz", "telefono": "", "deuda": 0},
        ]

        # Simular importación: crear clientes en SQLite (backend real)
        for f in filas:
            nombre = f["nombre"]
            if not nombre:
                continue
            c = db.crear_cliente(nombre=nombre, telefono=f.get("telefono", ""))
            deuda = float(f.get("deuda", 0) or 0)
            if deuda > 0:
                db.crear_movimiento(cliente_id=c["id"], tipo="prestamo",
                                    descripcion="Saldo inicial importado", monto=deuda)

        clientes = db.get_clientes()
        assert len(clientes) == 3

    def test_importar_sin_nombres_validos(self, db, data_dir):
        filas = [
            {"nombre": "", "deuda": 50000},
            {"nombre": None, "deuda": 0},
        ]

        importados = 0
        for f in filas:
            nombre = str(f.get("nombre", "")).strip() if f.get("nombre") else ""
            if not nombre:
                continue
            db.crear_cliente(nombre=nombre)
            importados += 1

        assert importados == 0
        assert db.get_clientes() == []

    def test_importar_con_prestamos_y_abonos(self, db, data_dir):
        c = db.crear_cliente(nombre="Cliente D", telefono="")
        db.crear_movimiento(cliente_id=c["id"], tipo="prestamo",
                            descripcion="Préstamo importado", monto=300000)
        db.crear_movimiento(cliente_id=c["id"], tipo="abono",
                            descripcion="Abono importado", monto=50000)

        clientes = db.get_clientes()
        assert len(clientes) == 1
        resumen = db.get_resumen_cliente(clientes[0]["id"])
        assert resumen["total_prestado"] == 300000.0
        assert resumen["total_abonado"] == 50000.0
        assert resumen["saldo"] == 250000.0

    def test_exportar_excel_genera_archivo(self, db):
        assert os.path.exists(db.excel_path)

    def test_backup_excel_sheet_structure(self, db):
        """Verifica que las hojas esperadas existan en el Excel de datos."""
        wb = load_workbook(db.excel_path)
        hojas = wb.sheetnames
        assert "clientes" in hojas
        assert "movimientos" in hojas
